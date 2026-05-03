from __future__ import annotations

from pathlib import Path

from docx import Document

from server.docx_io import ParsedDoc, Paragraph, parse_docx, write_docx


def test_parse_returns_heading_and_paragraphs(fixtures_dir: Path) -> None:
    parsed = parse_docx(fixtures_dir / "sample-en-folder" / "ch01.docx")
    assert isinstance(parsed, ParsedDoc)
    assert parsed.paragraphs[0] == Paragraph(style="heading-1", text="Chapter 1")
    assert parsed.paragraphs[1].style == "normal"
    assert parsed.paragraphs[1].text.startswith("The morning when")


def test_parse_encodes_italic_with_single_asterisks(fixtures_dir: Path) -> None:
    parsed = parse_docx(fixtures_dir / "sample-en-folder" / "ch01.docx")
    cold_para = next(p for p in parsed.paragraphs if "cold" in p.text)
    assert "*cold*" in cold_para.text


def test_parse_encodes_bold_with_double_asterisks(fixtures_dir: Path) -> None:
    parsed = parse_docx(fixtures_dir / "sample-en-folder" / "ch02.docx")
    letter_para = next(p for p in parsed.paragraphs if "letter" in p.text)
    assert "**letter**" in letter_para.text


def test_parse_preserves_unicode(fixtures_dir: Path) -> None:
    parsed = parse_docx(fixtures_dir / "sample-fr-folder" / "ch01.docx")
    assert any("« Pourquoi moi ? »" in p.text for p in parsed.paragraphs)


def test_write_then_parse_round_trips(tmp_path: Path, fixtures_dir: Path) -> None:
    parsed = parse_docx(fixtures_dir / "sample-en-folder" / "ch01.docx")
    out = tmp_path / "out.docx"
    write_docx(parsed, out)
    re_parsed = parse_docx(out)
    assert re_parsed.paragraphs == parsed.paragraphs


def test_write_emits_heading_at_correct_level(tmp_path: Path) -> None:
    parsed = ParsedDoc(paragraphs=[
        Paragraph(style="heading-1", text="Chapter 9"),
        Paragraph(style="normal", text="Body."),
    ])
    out = tmp_path / "out.docx"
    write_docx(parsed, out)
    doc = Document(str(out))
    assert doc.paragraphs[0].style.name == "Heading 1"
    assert doc.paragraphs[0].text == "Chapter 9"


def test_write_emits_inline_italic_and_bold(tmp_path: Path) -> None:
    parsed = ParsedDoc(paragraphs=[
        Paragraph(style="normal", text="Plain *italic* and **bold** end."),
    ])
    out = tmp_path / "out.docx"
    write_docx(parsed, out)
    re_parsed = parse_docx(out)
    assert re_parsed.paragraphs == parsed.paragraphs


def test_write_handles_unicode(tmp_path: Path) -> None:
    parsed = ParsedDoc(paragraphs=[
        Paragraph(style="normal", text="« Pourquoi moi ? » pensa-t-elle."),
    ])
    out = tmp_path / "out.docx"
    write_docx(parsed, out)
    re_parsed = parse_docx(out)
    assert re_parsed.paragraphs == parsed.paragraphs
