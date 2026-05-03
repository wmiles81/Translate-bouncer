from __future__ import annotations

from pathlib import Path

import pytest

from server.chapter_split import (
    ChapterSplitError,
    SplitChapter,
    split_by_heading_style,
    split_by_pattern,
    split_docx,
)
from server.docx_io import parse_docx


def test_split_by_heading_style_returns_three_chapters(fixtures_dir: Path) -> None:
    parsed = parse_docx(fixtures_dir / "sample-en-whole.docx")
    chapters = split_by_heading_style(parsed, heading_style="heading-1")
    assert len(chapters) == 3
    assert [c.title for c in chapters] == ["Chapter 1", "Chapter 2", "Chapter 3"]
    # Each chapter retains its heading paragraph plus its body.
    assert chapters[0].doc.paragraphs[0].style == "heading-1"
    assert any("morning" in p.text for p in chapters[0].doc.paragraphs)


def test_split_by_pattern_works_when_no_headings(tmp_path: Path) -> None:
    # Construct a doc with no heading styles, just plain paragraphs whose
    # first lines match a chapter pattern.
    from docx import Document
    d = Document()
    d.add_paragraph("Chapter 1")
    d.add_paragraph("Body of one.")
    d.add_paragraph("Chapter 2")
    d.add_paragraph("Body of two.")
    p = tmp_path / "no_headings.docx"
    d.save(str(p))
    parsed = parse_docx(p)
    chapters = split_by_pattern(parsed, patterns=[r"^Chapter\s+\d+$"])
    assert [c.title for c in chapters] == ["Chapter 1", "Chapter 2"]


def test_split_docx_prefers_heading_style(fixtures_dir: Path) -> None:
    chapters = split_docx(
        fixtures_dir / "sample-en-whole.docx",
        heading_style="heading-1",
        fallback_patterns=[r"^Chapter\s+\d+"],
    )
    assert len(chapters) == 3


def test_split_docx_falls_back_to_patterns_when_headings_yield_one(tmp_path: Path) -> None:
    from docx import Document
    d = Document()
    d.add_paragraph("Chapter 1")  # not a heading style
    d.add_paragraph("Body one.")
    d.add_paragraph("Chapter 2")
    d.add_paragraph("Body two.")
    p = tmp_path / "patterned.docx"
    d.save(str(p))
    chapters = split_docx(
        p,
        heading_style="heading-1",
        fallback_patterns=[r"^Chapter\s+\d+"],
    )
    assert len(chapters) == 2


def test_split_docx_raises_when_unable_to_split(tmp_path: Path) -> None:
    from docx import Document
    d = Document()
    d.add_paragraph("Just a single block of text.")
    p = tmp_path / "unsplittable.docx"
    d.save(str(p))
    with pytest.raises(ChapterSplitError):
        split_docx(
            p,
            heading_style="heading-1",
            fallback_patterns=[r"^Chapter\s+\d+"],
        )
