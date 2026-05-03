"""Generate test fixture .docx files. Run once; fixtures committed to git.

Usage:
    python tests/server/fixtures/_make_fixtures.py
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt

HERE = Path(__file__).parent

EN_CHAPTERS = [
    ("Chapter 1", [
        ("normal", "The morning when it all began, it was still raining."),
        ("normal", 'She thought, "Why me?"'),
        ("normal", "The window was *cold* under her hand."),
    ]),
    ("Chapter 2", [
        ("normal", "By noon the rain had stopped."),
        ("normal", "He carried the **letter** to the door."),
    ]),
    ("Chapter 3", [
        ("normal", "Three days passed without a word."),
        ("normal", "She decided to leave."),
    ]),
]

FR_CHAPTERS = [
    ("Chapitre 1", [
        ("normal", "Le matin où tout commença, il pleuvait encore."),
        ("normal", 'Elle pensa : « Pourquoi moi ? »'),
        ("normal", "La fenêtre était *froide* sous sa main."),
    ]),
    ("Chapitre 2", [
        ("normal", "À midi la pluie avait cessé."),
        ("normal", "Il porta la **lettre** à la porte."),
    ]),
    ("Chapitre 3", [
        ("normal", "Trois jours passèrent sans un mot."),
        ("normal", "Elle décida de partir."),
    ]),
]


def _add_paragraph(doc, style_name: str, text: str) -> None:
    if style_name.startswith("heading"):
        level = int(style_name.split("-")[1])
        p = doc.add_heading(text, level=level)
    else:
        p = doc.add_paragraph()
        # Render *italic*/**bold** markers as runs.
        i = 0
        while i < len(text):
            if text[i:i+2] == "**":
                end = text.find("**", i + 2)
                if end == -1:
                    p.add_run(text[i:])
                    break
                run = p.add_run(text[i+2:end])
                run.bold = True
                i = end + 2
            elif text[i] == "*":
                end = text.find("*", i + 1)
                if end == -1:
                    p.add_run(text[i:])
                    break
                run = p.add_run(text[i+1:end])
                run.italic = True
                i = end + 1
            else:
                # Take a chunk up to the next marker.
                next_marker = min(
                    [j for j in (text.find("*", i), text.find("**", i)) if j != -1] or [len(text)]
                )
                p.add_run(text[i:next_marker])
                i = next_marker


def write_folder(out_dir: Path, chapters):
    out_dir.mkdir(parents=True, exist_ok=True)
    for n, (heading, paras) in enumerate(chapters, start=1):
        doc = Document()
        _add_paragraph(doc, "heading-1", heading)
        for style, text in paras:
            _add_paragraph(doc, style, text)
        doc.save(str(out_dir / f"ch{n:02d}.docx"))


def write_whole(out_path: Path, chapters):
    doc = Document()
    for heading, paras in chapters:
        _add_paragraph(doc, "heading-1", heading)
        for style, text in paras:
            _add_paragraph(doc, style, text)
    doc.save(str(out_path))


if __name__ == "__main__":
    write_folder(HERE / "sample-en-folder", EN_CHAPTERS)
    write_folder(HERE / "sample-fr-folder", FR_CHAPTERS)
    write_whole(HERE / "sample-en-whole.docx", EN_CHAPTERS)
    write_whole(HERE / "sample-fr-whole.docx", FR_CHAPTERS)
    print("Fixtures written to:", HERE)
